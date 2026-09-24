"""Build the Excel workbook from data/*.csv and config.yaml.

Every input value comes from a CSV row. Every derived value is an Excel formula,
so the workbook recalculates when an input changes. The script also writes
output/cell_map.json, which maps model output names to cells, for parity_check.py.

Run:  python scripts/build_workbook.py
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import load_config, load_data, path, validate  # noqa: E402

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HDR = Font(name=FONT, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
SFILL = PatternFill("solid", fgColor="D9E1F2")
YFILL = PatternFill("solid", fgColor="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top")
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
MULT = '0.0"x"'
DATEF = "yyyy-mm-dd"
FMT = {"num": NUM, "num1": NUM1, "pct": PCT}

CELLMAP: dict[str, str] = {}


def reg(key, sheet, ref):
    CELLMAP[key] = f"{sheet}!{ref}"


def hdr(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def sec(ws, row, text, ncol):
    for i in range(1, ncol + 1):
        ws.cell(row=row, column=i).fill = SFILL
    ws.cell(row=row, column=1, value=text).font = BOLD


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, wrap=False):
    c = ws[ref]
    c.value = val
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    return c


def font_for(v):
    if isinstance(v, str) and v.startswith("="):
        return GREEN if "!" in v else BLACK
    return BLUE


def widths(ws, w):
    for k, v in w.items():
        ws.column_dimensions[k].width = v


def cval(x):
    """CSV value -> cell value (number if numeric, date if ISO date, else text)."""
    if x is None:
        return None
    s = str(x).strip()
    if s == "":
        return None
    try:
        return float(s) if "." in s else int(s)
    except ValueError:
        pass
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return dt.date.fromisoformat(s)
    return s


def table_sheet(wb, title, rows, cols, widths_map, blue_cols=(), header_row=1, note=None):
    ws = wb.create_sheet(title)
    start = header_row
    if note:
        put(ws, "A1", note, BOLD, wrap=True)
        ws.row_dimensions[1].height = 45
        start = 3
    hdr(ws, start, cols)
    keys = list(rows[0].keys()) if rows else []
    for r, row in enumerate(rows, start=start + 1):
        for c, k in enumerate(keys, start=1):
            v = cval(row[k])
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BLUE if k in blue_cols else BLACK
            cell.alignment = WRAP
            if isinstance(v, dt.date):
                cell.number_format = DATEF
    widths(ws, widths_map)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    return ws


def build(cfg, data):
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    blocks = cfg["blocks"]
    quarters = cfg["quarters"]
    qlab = [q["label"] for q in quarters]
    nq = len(qlab)
    qcols = [get_column_letter(2 + i) for i in range(nq)]
    lastc = qcols[-1]
    A = data["assumptions"]

    # ---------------- README
    ws = wb.active
    ws.title = "README"
    lines = [
        (f"{cfg['company']['name']} neocloud model — built {dt.date.today().isoformat()} from data/*.csv", TITLE),
        ("Do not edit numbers in this workbook. Edit the CSV files in data/, then run: python scripts/run_all.py", BOLD),
        ("", None),
        ("Color legend", BOLD),
        ("Blue text = input from a data file (each has a source in the Sources tab).", BLUE),
        ("Black text = formula in this sheet.", BLACK),
        ("Green text = link to another tab.", GREEN),
        ("Yellow fill = analyst assumption.", None),
        ("", None),
        ("Tabs: Sources, Reported_IS_CF, Quarterly, Balance_Sheet, KPIs, Assumptions, Analysis, Claim_Log, Reconciliation, Checks, Change_Log.", None),
        ("Verification status of each source and number: reports/verification.csv (run scripts/verify.py after scripts/fetch_sources.py).", None),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = f or BLACK
        if t.startswith("Yellow"):
            c.fill = YFILL
    ws.column_dimensions["A"].width = 140

    # ---------------- Sources
    table_sheet(wb, "Sources", data["sources"],
                ["ID", "Document", "Publisher", "Doc date", "Type", "Used for", "URL", "Accessed", "Confidence", "Verified vs primary?"],
                {"A": 6, "B": 48, "C": 16, "D": 11, "E": 11, "F": 45, "G": 60, "H": 11, "I": 11, "J": 22},
                blue_cols=("doc_date", "confidence", "verified_vs_primary"))

    # ---------------- Reported_IS_CF
    ws = wb.create_sheet("Reported_IS_CF")
    put(ws, "A1", "Reported lines ($ millions, as reported). Interest expense shown as a positive expense. Source: data/reported.csv", BOLD)
    hdr(ws, 3, ["Line item ($M)"] + [b["label"] for b in blocks])
    put(ws, "A4", "Period start", BOLD)
    put(ws, "A5", "Period end", BOLD)
    put(ws, "A6", "Source doc", BOLD)
    bcol = {}
    for i, b in enumerate(blocks):
        col = get_column_letter(2 + i)
        bcol[b["label"]] = col
        put(ws, f"{col}4", b["start"], BLUE, DATEF)
        put(ws, f"{col}5", b["end"], BLUE, DATEF)
        docs = {r["doc_id"] for r in data["reported_rows"] if r["block"] == b["label"]}
        put(ws, f"{col}6", "/".join(sorted(docs)), BLUE)
    rep_row = {}
    for i, item in enumerate(data["items"]):
        r = 7 + i
        rep_row[item] = r
        put(ws, f"A{r}", item)
        for b in blocks:
            put(ws, f"{bcol[b['label']]}{r}", data["reported"][(item, b["label"])], BLUE, NUM)
    widths(ws, {"A": 40, **{get_column_letter(2 + i): 13 for i in range(len(blocks))}})
    ws.freeze_panes = "B7"
    R = "Reported_IS_CF"

    # ---------------- Quarterly
    ws = wb.create_sheet("Quarterly")
    put(ws, "A1", "Standalone quarters ($ millions). Formulas on Reported_IS_CF; derivation rules in config.yaml.", BOLD)
    hdr(ws, 3, ["Item"] + qlab)
    for lab, rr in (("Quarter start", 4), ("Quarter end", 5), ("Days in quarter", 6), ("Derivation", 7)):
        put(ws, f"A{rr}", lab, BOLD)
    for q, col in zip(quarters, qcols):
        put(ws, f"{col}4", q["start"], BLUE, DATEF)
        put(ws, f"{col}5", q["end"], BLUE, DATEF)
        put(ws, f"{col}6", f"={col}5-{col}4+1", BLACK, NUM)
        put(ws, f"{col}7", " + ".join(q["plus"]) + ("" if not q["minus"] else " − " + " − ".join(q["minus"])))
    q_row = {}
    for i, item in enumerate(data["items"]):
        r = 8 + i
        q_row[item] = r
        put(ws, f"A{r}", item)
        rr = rep_row[item]
        for q, col in zip(quarters, qcols):
            fx = "=" + "+".join(f"{R}!{bcol[b]}{rr}" for b in q["plus"]) + "".join(f"-{R}!{bcol[b]}{rr}" for b in q["minus"])
            put(ws, f"{col}{r}", fx, GREEN, NUM)
            reg(f"quarterly|{item}|{q['label']}", "Quarterly", f"{col}{r}")
    r0 = 8 + len(data["items"]) + 1
    sec(ws, r0, "Ratios", nq + 1)
    rev = q_row["Revenue"]
    ratio_defs = [
        ("growth_qoq", "Revenue growth, QoQ", PCT),
        ("growth_yoy", "Revenue growth, YoY", PCT),
        ("adj_ebitda_margin", "Adjusted EBITDA margin", PCT),
        ("op_margin", "Operating margin (GAAP)", PCT),
        ("da_pct_rev", "D&A as % of revenue", PCT),
        ("interest_pct_rev", "Interest expense as % of revenue", PCT),
        ("ebitda_interest_cover", "Adj. EBITDA / interest expense (x)", MULT),
        ("capex_to_rev", "Cash capex / revenue (x)", MULT),
    ]
    Arow = {}  # assumption key -> row, filled below; yoy formulas reference Assumptions
    assumption_keys = list(A.keys())
    for k, key in enumerate(assumption_keys):
        Arow[key] = 4 + k
    for k, (key, name, fmt) in enumerate(ratio_defs):
        r = r0 + 1 + k
        put(ws, f"A{r}", name)
        for j, (q, col) in enumerate(zip(qlab, qcols)):
            prev = qcols[j - 1] if j > 0 else None
            fx = None
            if key == "growth_qoq":
                fx = None if prev is None else f"={col}{rev}/{prev}{rev}-1"
            elif key == "growth_yoy":
                y = cfg["yoy"].get(q)
                if y:
                    if "assumption" in y:
                        base = f"Assumptions!$C${Arow[y['assumption']]}"
                    elif "block" in y:
                        base = f"{R}!{bcol[y['block']]}{rep_row['Revenue']}"
                    else:
                        base = f"{qcols[qlab.index(y['quarter'])]}{rev}"
                    fx = f"={col}{rev}/{base}-1"
            elif key == "adj_ebitda_margin":
                fx = f"={col}{q_row['Adjusted EBITDA']}/{col}{rev}"
            elif key == "op_margin":
                fx = f"={col}{q_row['Operating income (loss)']}/{col}{rev}"
            elif key == "da_pct_rev":
                fx = f"={col}{q_row['Depreciation and amortization']}/{col}{rev}"
            elif key == "interest_pct_rev":
                fx = f"={col}{q_row['Interest expense, net']}/{col}{rev}"
            elif key == "ebitda_interest_cover":
                fx = f"={col}{q_row['Adjusted EBITDA']}/{col}{q_row['Interest expense, net']}"
            elif key == "capex_to_rev":
                fx = f"={col}{q_row['Purchase of PP&E (cash capex)']}/{col}{rev}"
            if fx is None:
                put(ws, f"{col}{r}", "n/a")
            else:
                put(ws, f"{col}{r}", fx, BLACK, fmt)
                reg(f"metric|{key}|{q}", "Quarterly", f"{col}{r}")
    widths(ws, {"A": 40, **{c: 13 for c in qcols}})
    ws.freeze_panes = "B8"

    # ---------------- Balance_Sheet
    ws = wb.create_sheet("Balance_Sheet")
    put(ws, "A1", "Balance sheet ($ millions, as reported). Source: data/balance_sheet.csv", BOLD)
    dates = data["bs_dates"]
    dcols = [get_column_letter(2 + i) for i in range(len(dates))]
    hdr(ws, 3, ["Item"] + dates)
    put(ws, "A4", "Source doc", BOLD)
    for d, col in zip(dates, dcols):
        docs = {r["doc_id"] for r in data["bs_rows"] if r["as_of"] == d}
        put(ws, f"{col}4", "/".join(sorted(docs)), BLUE)
    bs_row = {}
    for i, item in enumerate(data["bs_items"]):
        r = 5 + i
        bs_row[item] = r
        put(ws, f"A{r}", item)
        for d, col in zip(dates, dcols):
            put(ws, f"{col}{r}", data["bs"][(item, d)], BLUE, NUM)
    r = 5 + len(data["bs_items"]) + 1
    sec(ws, r, "Derived", len(dates) + 1)
    derived = [
        ("Total debt", ["Debt, current (recourse + non-recourse)", "Debt, non-current (recourse + non-recourse)"]),
        ("Total deferred revenue", ["Deferred revenue, current", "Deferred revenue, non-current"]),
        ("Total operating lease liabilities", ["Operating lease liabilities, current", "Operating lease liabilities, non-current"]),
        ("Total cash incl. restricted", ["Cash and cash equivalents", "Restricted cash, current", "Restricted cash, non-current"]),
    ]
    for k, (name, parts) in enumerate(derived):
        rr = r + 1 + k
        bs_row[name] = rr
        put(ws, f"A{rr}", name)
        for d, col in zip(dates, dcols):
            put(ws, f"{col}{rr}", "=" + "+".join(f"{col}{bs_row[p]}" for p in parts), BLACK, NUM)
            if name == "Total debt":
                reg(f"metric|total_debt|{d}", "Balance_Sheet", f"{col}{rr}")
    rr = r + 1 + len(derived)
    bs_row["Net debt (excl. leases)"] = rr
    put(ws, f"A{rr}", "Net debt (excl. leases)")
    put(ws, f"A{rr+1}", "Net debt incl. operating leases")
    for col in dcols:
        put(ws, f"{col}{rr}", f"={col}{bs_row['Total debt']}-{col}{bs_row['Total cash incl. restricted']}", BLACK, NUM)
        put(ws, f"{col}{rr+1}", f"={col}{rr}+{col}{bs_row['Total operating lease liabilities']}", BLACK, NUM)
    widths(ws, {"A": 45, **{c: 13 for c in dcols}})
    BSn = "Balance_Sheet"
    dcol = dict(zip(dates, dcols))

    # ---------------- KPIs
    ws = wb.create_sheet("KPIs")
    put(ws, "A1", "Operating KPIs at quarter end. Source: data/kpis.csv (value, qualifier, source, evidence text).", BOLD)
    hdr(ws, 3, ["Item"] + qlab)
    labels = {"active_mw": ("Active power, end of quarter (MW)", NUM), "contracted_gw": ("Contracted power, end of quarter (GW)", NUM1),
              "backlog_busd": ("Revenue backlog, end of quarter ($B)", NUM1), "data_centers": ("Active data centers (count)", NUM)}
    kp_row = {}
    r = 4
    for m in data["kpi_metrics"]:
        name, fmt = labels.get(m, (m, NUM1))
        kp_row[m] = r
        put(ws, f"A{r}", name, BOLD)
        put(ws, f"A{r+1}", "  qualifier")
        put(ws, f"A{r+2}", "  source")
        for q, col in zip(qlab, qcols):
            row = data["kpis"].get((m, q))
            if row:
                put(ws, f"{col}{r}", float(row["value"]), BLUE, fmt)
                put(ws, f"{col}{r+1}", row["qualifier"])
                put(ws, f"{col}{r+2}", row["doc_id"])
        r += 3
    widths(ws, {"A": 40, **{c: 16 for c in qcols}})
    ACT, CON, BKL = kp_row["active_mw"], kp_row["contracted_gw"], kp_row["backlog_busd"]

    # ---------------- Assumptions
    ws = wb.create_sheet("Assumptions")
    put(ws, "A1", "Assumptions and guidance inputs. Source: data/assumptions.csv", BOLD)
    hdr(ws, 3, ["Item", "Key", "Value", "Unit", "Basis", "Source", "Confidence"])
    for key in assumption_keys:
        a = A[key]
        r = Arow[key]
        put(ws, f"A{r}", a["label"], wrap=True)
        put(ws, f"B{r}", key)
        put(ws, f"C{r}", a["value"], BLUE, NUM2 if a["unit"] == "kW/GPU" else NUM, YFILL if a["is_key"] == "1" else None)
        put(ws, f"D{r}", a["unit"])
        put(ws, f"E{r}", a["basis"], wrap=True)
        put(ws, f"F{r}", a["doc_id"])
        put(ws, f"G{r}", a["confidence"])
    r_q3 = 4 + len(assumption_keys)
    put(ws, f"A{r_q3}", "Q3'26 end active power (linear path from last actual to YE26 target)")
    put(ws, f"B{r_q3}", "q3_26_active_mw_linear")
    put(ws, f"C{r_q3}", f"=(KPIs!{lastc}{ACT}+C{Arow['ye26_active_mw']})/2", GREEN, NUM, YFILL)
    put(ws, f"D{r_q3}", "MW")
    put(ws, f"E{r_q3}", "Analyst assumption: straight line.")
    put(ws, f"G{r_q3}", "Low")
    widths(ws, {"A": 60, "B": 22, "C": 12, "D": 9, "E": 60, "F": 9, "G": 11})
    KW = f"Assumptions!$C${Arow['kw_per_gpu']}"
    HPD = f"Assumptions!$C${Arow['hours_per_day']}"

    # ---------------- Analysis
    ws = wb.create_sheet("Analysis")
    put(ws, "A1", "Analysis (formulas only). Columns match Quarterly.", BOLD)
    hdr(ws, 3, ["Metric"] + qlab)

    def arow(r, key, name, fn, fmt, first_na=False):
        put(ws, f"A{r}", name)
        for j, (q, col) in enumerate(zip(qlab, qcols)):
            prev = qcols[j - 1] if j > 0 else None
            if first_na and j == 0:
                put(ws, f"{col}{r}", "n/a")
                continue
            fx = fn(col, prev)
            put(ws, f"{col}{r}", fx, font_for(fx) if fx.startswith("=Quarterly") or fx.startswith("=KPIs") else BLACK, fmt)
            if key:
                reg(f"metric|{key}|{q}", "Analysis", f"{col}{r}")

    sec(ws, 4, "A. Capacity economics", nq + 1)
    arow(5, "avg_active_mw", "Average active power in quarter (MW)", lambda c, p: f"=(KPIs!{p}{ACT}+KPIs!{c}{ACT})/2", NUM, True)
    arow(6, None, "Revenue ($M)", lambda c, p: f"=Quarterly!{c}{rev}", NUM)
    arow(7, "rev_per_mw", "Annualized revenue per average active MW ($M per MW-year)", lambda c, p: f"={c}6*365/Quarterly!{c}6/{c}5", NUM2, True)
    arow(8, "ebitda_per_mw", "Annualized adj. EBITDA per average active MW ($M per MW-year)", lambda c, p: f"=Quarterly!{c}{q_row['Adjusted EBITDA']}*365/Quarterly!{c}6/{c}5", NUM2, True)
    arow(9, "gpus_k", "Implied installed GPUs (thousands) = MW / kW per GPU", lambda c, p: f"={c}5/{KW}", NUM1, True)
    arow(10, "gpu_hours_m", "Implied installed GPU-hours in quarter (millions)", lambda c, p: f"={c}9*Quarterly!{c}6*{HPD}/1000", NUM1, True)
    arow(11, "rev_per_gpu_hour", "Implied revenue per installed GPU-hour ($)", lambda c, p: f"={c}6/{c}10", USD2, True)
    put(ws, "A12", "Units: $M ÷ million GPU-hours = $ per GPU-hour. Counts all active capacity, not only billed hours: a lower limit on the realized price.", wrap=True)
    sec(ws, 14, "B. Backlog", nq + 1)
    arow(15, "backlog", "Revenue backlog, end ($B)", lambda c, p: f"=KPIs!{c}{BKL}", NUM1)
    arow(16, "bookings", "Implied gross bookings in quarter ($B) = change in backlog + revenue", lambda c, p: f"={c}15-{p}15+Quarterly!{c}{rev}/1000", NUM1, True)
    arow(17, "backlog_years", "Backlog / annualized revenue (years)", lambda c, p: f"={c}15*1000/(Quarterly!{c}{rev}*365/Quarterly!{c}6)", NUM1)
    sec(ws, 19, "C. Capacity pipeline", nq + 1)
    arow(20, "contracted_gw", "Contracted power, end (GW)", lambda c, p: f"=KPIs!{c}{CON}", NUM2)
    arow(21, "active_gw", "Active power, end (GW)", lambda c, p: f"=KPIs!{c}{ACT}/1000", NUM2)
    arow(22, "pipeline_gw", "Contracted but not yet active (GW)", lambda c, p: f"={c}20-{c}21", NUM2)
    arow(23, "added_mw", "Active power added in quarter (MW)", lambda c, p: f"=KPIs!{c}{ACT}-KPIs!{p}{ACT}", NUM, True)
    arow(24, "quarters_to_activate", "Quarters to activate the pipeline at this quarter's add rate", lambda c, p: f"=IF({c}23<=0,\"n/a\",{c}22*1000/{c}23)", NUM1, True)
    arow(25, "active_share", "Active share of contracted power", lambda c, p: f"={c}21/{c}20", PCT)
    arow(26, "backlog_per_gw", "Backlog per GW of contracted power ($B/GW)", lambda c, p: f"={c}15/{c}20", NUM1)

    sec(ws, 28, "D. Implied depreciation rate and interest rate (period blocks)", nq + 1)
    rps = cfg["rate_periods"]
    hdr(ws, 29, ["Metric"] + [rp["label"] for rp in rps])
    ppe, td = bs_row["Property and equipment, net"], bs_row["Total debt"]
    labels_d = ["Days in period", "Annualization factor", "D&A ($M)", "Average PP&E, net ($M)", "Implied annual depreciation rate on net PP&E",
                "Implied average life (years) = 1 / rate", "Interest expense, net ($M)", "Average total debt ($M)", "Implied annual interest rate on average debt"]
    for k, lab in enumerate(labels_d):
        put(ws, f"A{30+k}", lab)
    for j, rp in enumerate(rps):
        c = get_column_letter(2 + j)
        s, e = rp["bs_start"].isoformat(), rp["bs_end"].isoformat()
        put(ws, f"{c}30", f"=DATE({rp['end'].year},{rp['end'].month},{rp['end'].day})-DATE({rp['start'].year},{rp['start'].month},{rp['start'].day})+1", BLACK, NUM)
        put(ws, f"{c}31", f"=365/{c}30", BLACK, NUM2)
        put(ws, f"{c}32", f"={R}!{bcol[rp['block']]}{rep_row['Depreciation and amortization']}", GREEN, NUM)
        put(ws, f"{c}33", f"=({BSn}!{dcol[s]}{ppe}+{BSn}!{dcol[e]}{ppe})/2", GREEN, NUM)
        put(ws, f"{c}34", f"={c}32*{c}31/{c}33", BLACK, PCT)
        put(ws, f"{c}35", f"=1/{c}34", BLACK, NUM1)
        put(ws, f"{c}36", f"={R}!{bcol[rp['block']]}{rep_row['Interest expense, net']}", GREEN, NUM)
        put(ws, f"{c}37", f"=({BSn}!{dcol[s]}{td}+{BSn}!{dcol[e]}{td})/2", GREEN, NUM)
        put(ws, f"{c}38", f"={c}36*{c}31/{c}37", BLACK, PCT)
        reg(f"metric|dep_rate|{rp['label']}", "Analysis", f"{c}34")
        reg(f"metric|implied_life|{rp['label']}", "Analysis", f"{c}35")
        reg(f"metric|interest_rate|{rp['label']}", "Analysis", f"{c}38")
    put(ws, "A39", "Caveats: net PP&E includes construction in progress, so the implied life is too long; interest is net of capitalized interest, so the implied rate is too low.", wrap=True)

    sec(ws, 41, "E. FY26 guidance check", nq + 1)
    hdr(ws, 42, ["Metric", "Low", "High"])
    lab_e = ["H1'26 revenue, reported ($M)", "FY26 revenue guidance ($M)", "Implied H2'26 revenue ($M)", "Q3'26 revenue guidance ($M)",
             "Implied Q4'26 revenue ($M)", "Implied Q4'26 growth vs Q3'26", "Implied Q4'26 average active power (MW), linear path",
             "Implied Q4'26 annualized revenue per active MW ($M per MW-year)", f"{qlab[-1]} actual, same metric", "Change in revenue per MW needed"]
    for k, lab in enumerate(lab_e):
        put(ws, f"A{43+k}", lab)
    h1_col = bcol[cfg.get("h1_block", "6M'26")]
    for c, side in (("B", "low"), ("C", "high")):
        put(ws, f"{c}43", f"={R}!{h1_col}{rep_row['Revenue']}", GREEN, NUM)
        put(ws, f"{c}44", f"=Assumptions!C{Arow['fy26_rev_' + side]}", GREEN, NUM)
        put(ws, f"{c}45", f"={c}44-{c}43", BLACK, NUM)
        put(ws, f"{c}46", f"=Assumptions!C{Arow['q3_26_rev_' + side]}", GREEN, NUM)
        put(ws, f"{c}47", f"={c}45-{c}46", BLACK, NUM)
        put(ws, f"{c}48", f"={c}47/{c}46-1", BLACK, PCT)
        put(ws, f"{c}49", f"=(Assumptions!C{r_q3}+Assumptions!C{Arow['ye26_active_mw']})/2", GREEN, NUM)
        put(ws, f"{c}50", f"={c}47*365/92/{c}49", BLACK, NUM2)
        put(ws, f"{c}51", f"={lastc}7", BLACK, NUM2)
        put(ws, f"{c}52", f"={c}50/{c}51-1", BLACK, PCT)
        for key, rr in (("g_h2_rev", 45), ("g_q4_rev", 47), ("g_q4_growth", 48), ("g_q4_avg_mw", 49), ("g_q4_rev_per_mw", 50), ("g_change_needed", 52)):
            reg(f"metric|{key}|{side}", "Analysis", f"{c}{rr}")

    sec(ws, 54, f"F. Sensitivity: {qlab[-1]} implied revenue per installed GPU-hour vs kW per GPU", nq + 1)
    hdr(ws, 55, ["kW per GPU", "$ per GPU-hour"])
    for i, k in enumerate(cfg["sensitivity_kw_per_gpu"]):
        r = 56 + i
        put(ws, f"A{r}", k, BLUE, NUM1)
        put(ws, f"B{r}", f"=Quarterly!{lastc}{rev}/((${lastc}$5/A{r})*Quarterly!{lastc}6*24/1000)", BLACK, USD2)
        reg(f"metric|sens_rev_per_gpu_hour|{k}", "Analysis", f"B{r}")
    widths(ws, {"A": 62, **{c: 14 for c in qcols}})
    ws.freeze_panes = "B4"
    AN = {"bookings": 16, "active_share": 25}

    # ---------------- Claim_Log
    table_sheet(wb, "Claim_Log", data["claims"],
                ["ID", "Company", "Claim date", "Metric", "Value", "Unit", "Qualifier", "Class", "As-of / target date", "Outcome / status", "Source", "Confidence"],
                {"A": 6, "B": 11, "C": 11, "D": 34, "E": 9, "F": 6, "G": 14, "H": 12, "I": 16, "J": 44, "K": 8, "L": 11},
                blue_cols=("claim_date", "value", "as_of_or_target"),
                note="Capacity claims. Announced = plan or target, no binding contract stated. Contracted = power under lease/contract. "
                     "Active = energized and in service (may include unbilled capacity). Available = capacity the company says it can sell in a stated window.")

    # ---------------- Reconciliation
    ws = wb.create_sheet("Reconciliation")
    put(ws, "A1", "Conflicting figures and how each is resolved. Source: data/reconciliation.csv", BOLD)
    hdr(ws, 3, ["ID", "Topic", "Figure A", "Source A", "Figure B", "Source B", "Difference (A − B)", "Likely cause", "Status", "Next step"])

    def ref_formula(tok):
        s = str(tok).strip()
        m = re.fullmatch(r"\{(.+):(.+)\}", s)
        if not m:
            v = cval(s)
            return v
        name, per = m.group(1), m.group(2)
        col = qcols[qlab.index(per)]
        if name in q_row:
            return f"=Quarterly!{col}{q_row[name]}"
        if name in AN:
            return f"=Analysis!{col}{AN[name]}"
        raise KeyError(s)

    for r, row in enumerate(data["reconciliation"], start=4):
        fa, fb = ref_formula(row["figure_a"]), ref_formula(row["figure_b"])
        fmt = FMT.get(row["fmt"], NUM)
        vals = [row["rec_id"], row["topic"], fa, row["source_a"], fb, row["source_b"], f"=C{r}-E{r}", row["likely_cause"], row["status"], row["next_step"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = WRAP
            cell.font = font_for(v) if c in (3, 5) else BLACK
            if c in (3, 5, 7):
                cell.number_format = fmt
    widths(ws, {"A": 5, "B": 28, "C": 11, "D": 24, "E": 11, "F": 24, "G": 12, "H": 40, "I": 10, "J": 32})

    # ---------------- Checks
    ws = wb.create_sheet("Checks")
    put(ws, "A1", "Checks (TRUE = pass). Cross-source checks come from data/xchecks.csv.", BOLD)
    hdr(ws, 3, ["ID", "Check", "Result", "Tolerance", "Model value", "Reference value", "Reference source"])
    checks = []
    for q, col in zip(qlab, qcols):
        comp = "+".join(f"Quarterly!{col}{q_row[n]}" for n in ["Cost of revenue", "Technology and infrastructure", "Sales and marketing", "General and administrative"])
        checks.append((f"A-opex-{q}", f"Opex lines sum to total opex, {q}", f"={comp}", f"=Quarterly!{col}{q_row['Total operating expenses']}", 0.5, "Arithmetic"))
        checks.append((f"A-opinc-{q}", f"Revenue - opex = operating income, {q}", f"=Quarterly!{col}{rev}-Quarterly!{col}{q_row['Total operating expenses']}", f"=Quarterly!{col}{q_row['Operating income (loss)']}", 0.5, "Arithmetic"))
    for d, col in zip(dates, dcols):
        lhs = f"={BSn}!{col}{bs_row['Total liabilities']}+{BSn}!{col}{bs_row['Redeemable convertible preferred stock']}+{BSn}!{col}{bs_row[chr(84) + 'otal stockholders' + chr(39) + ' equity (deficit)']}"
        checks.append((f"A-bs-{d}", f"Liabilities + preferred + equity = assets, {d}", lhs, f"={BSn}!{col}{bs_row['Total assets']}", 0.5, "Arithmetic"))
    for blk in sorted({b["block"] for b in data["bridge"]}):
        extra = sum(float(b["value_musd"]) for b in data["bridge"] if b["block"] == blk)
        c = bcol[blk]
        lhs = f"={R}!{c}{rep_row['Net loss']}+{R}!{c}{rep_row['Depreciation and amortization']}+{R}!{c}{rep_row['Interest expense, net']}+{R}!{c}{rep_row['Stock-based compensation']}+({extra:g})"
        checks.append((f"A-bridge-{blk}", f"Adj. EBITDA bridge, {blk} (other bridge items from ebitda_bridge.csv)", lhs, f"={R}!{c}{rep_row['Adjusted EBITDA']}", 1, "ebitda_bridge.csv"))
    QR = 8 + len(data["items"]) + 1  # ratio section start in Quarterly
    ratio_row = {"growth_qoq": QR + 1, "growth_yoy": QR + 2}
    for x in data["xchecks"]:
        m = re.fullmatch(r"\{(.+):(.+)\}", x["model_ref"].strip())
        name, per = m.group(1), m.group(2)
        if name in q_row:
            mv = f"=Quarterly!{qcols[qlab.index(per)]}{q_row[name]}"
        elif name in ratio_row:
            mv = f"=Quarterly!{qcols[qlab.index(per)]}{ratio_row[name]}"
        elif name == "total_debt":
            mv = f"={BSn}!{dcol[per]}{bs_row['Total debt']}"
        else:
            raise KeyError(x["model_ref"])
        checks.append((x["check_id"], x["label"], mv, float(x["reference_value"]), float(x["tolerance"]), x["reference_source"]))
    for i, (cid, label, mv, rv, tol, src) in enumerate(checks, start=4):
        put(ws, f"A{i}", cid)
        put(ws, f"B{i}", label, wrap=True)
        put(ws, f"D{i}", tol, BLUE, NUM2)
        put(ws, f"E{i}", mv, GREEN, NUM2)
        put(ws, f"F{i}", rv, font_for(rv), NUM2)
        put(ws, f"C{i}", f"=ABS(E{i}-F{i})<=D{i}")
        put(ws, f"G{i}", src)
        reg(f"check|{cid}|", "Checks", f"C{i}")
    last = 3 + len(checks)
    put(ws, f"B{last+2}", "All checks pass?", BOLD)
    put(ws, f"C{last+2}", f"=COUNTIF(C4:C{last},FALSE)=0", BOLD)
    reg("check|ALL|", "Checks", f"C{last+2}")
    widths(ws, {"A": 16, "B": 60, "C": 9, "D": 10, "E": 13, "F": 14, "G": 22})

    # ---------------- Change_Log
    table_sheet(wb, "Change_Log", data["change_log"], ["Date", "Tab", "Item", "Old value", "New value", "Reason", "Source"],
                {"A": 11, "B": 13, "C": 24, "D": 10, "E": 12, "F": 60, "G": 12})
    return wb


def main() -> int:
    cfg, data = load_config(), load_data()
    problems = validate(data, cfg)
    if problems:
        print("DATA PROBLEMS:")
        for p in problems:
            print("  -", p)
        return 1
    wb = build(cfg, data)
    out = path(cfg, "output")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    with open(out.parent / "cell_map.json", "w", encoding="utf-8") as f:
        json.dump(CELLMAP, f, indent=1)
    print(f"workbook: {out} ({len(CELLMAP)} mapped outputs)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
